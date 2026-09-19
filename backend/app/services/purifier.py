import asyncio
import time
import uuid
from pathlib import Path
from typing import Any, Dict

from fastapi import UploadFile
from openpyxl import load_workbook

from app.core.config import settings
from app.services.components.docx_component import docx_component
from app.services.components.pdf_component import pdf_component
from app.services.components.text_component import text_component


class UnsupportedFileError(ValueError):
    """Raised when an uploaded file fails extension or size validation."""


EXCEL_EXTENSIONS = {
    ".xlsx",
    ".xlsm",
    ".xltx",
    ".xltm",
}

EXCEL_MAX_ROWS_PER_SHEET = 250
EXCEL_MAX_COLUMNS_PER_SHEET = 40
EXCEL_MAX_EXTRACTED_CHARS = 80000


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))

    return str(value).strip()


def _extract_excel_sync(file_path: Path) -> str:
    """
    Extract structured spreadsheet content for local LLM analysis.

    The original .xlsx/.xlsm file remains untouched on disk. Only a bounded,
    text representation is produced for the agent context.
    """
    try:
        workbook_values = load_workbook(
            filename=file_path,
            read_only=True,
            data_only=True,
        )

        try:
            workbook_formulas = load_workbook(
                filename=file_path,
                read_only=True,
                data_only=False,
            )
        except Exception:
            workbook_formulas = None

        parts = [
            f"EXCEL WORKBOOK: {file_path.name}",
            f"SHEETS: {len(workbook_values.sheetnames)}",
        ]

        for sheet_index, sheet_name in enumerate(
            workbook_values.sheetnames,
            start=1,
        ):
            value_sheet = workbook_values[sheet_name]
            formula_sheet = (
                workbook_formulas[sheet_name]
                if workbook_formulas is not None
                and sheet_name in workbook_formulas.sheetnames
                else None
            )

            parts.append("")
            parts.append(
                f"=== SHEET {sheet_index}: {sheet_name} ==="
            )

            row_count = 0
            non_empty_rows = 0

            for row in value_sheet.iter_rows(
                max_row=EXCEL_MAX_ROWS_PER_SHEET,
                max_col=EXCEL_MAX_COLUMNS_PER_SHEET,
            ):
                row_count += 1

                values = [
                    _cell_to_text(cell.value)
                    for cell in row
                ]

                if not any(values):
                    continue

                non_empty_rows += 1

                # Keep formula information when Excel has not cached a value.
                if formula_sheet is not None:
                    formula_row = next(
                        formula_sheet.iter_rows(
                            min_row=row_count,
                            max_row=row_count,
                            max_col=EXCEL_MAX_COLUMNS_PER_SHEET,
                        ),
                        (),
                    )

                    for index, value in enumerate(values):
                        if value:
                            continue

                        if index < len(formula_row):
                            formula = formula_row[index].value
                            if (
                                isinstance(formula, str)
                                and formula.startswith("=")
                            ):
                                values[index] = formula

                last_non_empty = 0
                for index, value in enumerate(values):
                    if value:
                        last_non_empty = index + 1

                values = values[:last_non_empty]

                parts.append(
                    f"Row {row_count}: "
                    + " | ".join(values)
                )

                if sum(len(part) + 1 for part in parts) >= EXCEL_MAX_EXTRACTED_CHARS:
                    parts.append(
                        "[Excel extraction truncated to keep agent context bounded.]"
                    )
                    break

            parts.append(
                f"Rows scanned: {row_count}; "
                f"non-empty rows: {non_empty_rows}"
            )

            if sum(len(part) + 1 for part in parts) >= EXCEL_MAX_EXTRACTED_CHARS:
                break

        return "\n".join(parts)[:EXCEL_MAX_EXTRACTED_CHARS]

    finally:
        try:
            workbook_values.close()
        except Exception:
            pass

        try:
            if workbook_formulas is not None:
                workbook_formulas.close()
        except Exception:
            pass


class DataPurifierService:
    """Routes each input to its dedicated component and normalizes the results into one purified payload."""

    def __init__(self, workspace_dir: Path):
        self.workspace_dir = workspace_dir
        self.workspace_dir.mkdir(parents=True, exist_ok=True)

    def sanitize_text(self, text: str) -> str:
        return text_component.extract(text)

    def cleanup_workspace(self, max_age_hours: int = None) -> int:
        """Deletes workspace files older than the retention window. Returns count removed."""
        max_age_hours = (
            max_age_hours
            if max_age_hours is not None
            else settings.WORKSPACE_RETENTION_HOURS
        )
        cutoff = time.time() - (max_age_hours * 3600)
        removed = 0

        for f in self.workspace_dir.iterdir():
            if f.is_file() and f.stat().st_mtime < cutoff:
                f.unlink(missing_ok=True)
                removed += 1

        return removed

    async def save_and_parse_file(
        self,
        file: UploadFile,
    ) -> Dict[str, Any]:
        """Saves file to local disk and routes it to the matching parser."""
        filename = file.filename or "uploaded_file"
        ext = Path(filename).suffix.lower()

        allowed_extensions = (
            set(settings.ALLOWED_EXTENSIONS)
            | EXCEL_EXTENSIONS
        )

        if ext not in allowed_extensions:
            raise UnsupportedFileError(
                f"Unsupported file type '{ext}' for '{filename}'"
            )

        file_id = f"doc_{uuid.uuid4().hex[:8]}"
        file_path = (
            self.workspace_dir
            / f"{file_id}_{Path(filename).name}"
        )

        contents = await file.read()

        max_bytes = (
            settings.MAX_UPLOAD_SIZE_MB
            * 1024
            * 1024
        )

        if len(contents) > max_bytes:
            raise UnsupportedFileError(
                f"File '{filename}' exceeds the "
                f"{settings.MAX_UPLOAD_SIZE_MB}MB upload limit"
            )

        with open(file_path, "wb") as output_file:
            output_file.write(contents)

        extracted_content = ""
        file_type = "unknown"

        # Extraction is CPU-bound and blocking: offload to a worker thread.
        if ext == ".pdf":
            file_type, extracted_content = await asyncio.to_thread(
                pdf_component.extract,
                file_path,
                self.workspace_dir,
            )

        elif ext in [".docx", ".doc"]:
            file_type = "word_document"
            extracted_content = await asyncio.to_thread(
                docx_component.extract,
                file_path,
            )

        elif ext in [
            ".png",
            ".jpg",
            ".jpeg",
            ".webp",
            ".bmp",
        ]:
            # Vision requests are handled directly by Qwen-VL.
            file_type = "image"
            extracted_content = ""

        elif ext in EXCEL_EXTENSIONS:
            file_type = "spreadsheet"
            extracted_content = await asyncio.to_thread(
                _extract_excel_sync,
                file_path,
            )

        elif ext in [
            ".txt",
            ".log",
            ".py",
            ".json",
            ".csv",
        ]:
            file_type = "plain_text"
            extracted_content = await asyncio.to_thread(
                file_path.read_text,
                encoding="utf-8",
                errors="ignore",
            )

        return {
            "file_id": file_id,
            "filename": filename,
            "local_path": file_path.relative_to(
                settings.BASE_DIR
            ).as_posix(),
            "file_type": file_type,
            "size_bytes": len(contents),
            "extracted_text": self.sanitize_text(
                extracted_content
            ),
        }
