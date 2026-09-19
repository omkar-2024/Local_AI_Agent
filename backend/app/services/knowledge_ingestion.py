from __future__ import annotations

import asyncio
from pathlib import Path
from typing import Any

import fitz
import docx

from app.core.config import settings
from app.services.components.image_component import image_component
from app.services.rag_engine import rag_engine
from app.services.vector_store import vector_store


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".log",
    ".csv",
    ".json",
    ".py",
    ".png",
    ".jpg",
    ".jpeg",
}


def _clean(text: str) -> str:
    return "\n".join(
        " ".join(line.split())
        for line in (text or "").splitlines()
        if line.strip()
    )


def _chunk_page(text: str, max_chars: int = 900, overlap: int = 120) -> list[str]:
    """Create compact, paragraph-aware chunks while preserving page locality."""
    text = _clean(text)
    if not text:
        return []

    paragraphs = [p for p in text.splitlines() if p.strip()]
    chunks: list[str] = []
    current = ""

    for paragraph in paragraphs:
        if len(paragraph) > max_chars:
            if current:
                chunks.append(current)
                current = ""

            start = 0
            while start < len(paragraph):
                end = start + max_chars
                piece = paragraph[start:end].strip()
                if piece:
                    chunks.append(piece)
                start = max(end - overlap, start + 1)
            continue

        candidate = (
            f"{current}\n{paragraph}".strip()
            if current
            else paragraph
        )

        if len(candidate) <= max_chars:
            current = candidate
            continue

        if current:
            chunks.append(current)
            tail = current[-overlap:]
            current = f"{tail}\n{paragraph}".strip()
        else:
            current = paragraph

    if current:
        chunks.append(current)

    return chunks


def _extract_file_sync(path: Path) -> list[dict[str, Any]]:
    ext = path.suffix.lower()

    if ext == ".pdf":
        doc = fitz.open(path)
        chunks: list[dict[str, Any]] = []

        try:
            for page_number, page in enumerate(doc, start=1):
                text = page.get_text("text").strip()

                if not text:
                    # Scanned page: OCR locally.
                    pix = page.get_pixmap(
                        dpi=130,
                        alpha=False,
                    )
                    image_path = (
                        settings.SANDBOX_DIR
                        / f"rag_ocr_{path.stem}_{page_number}.png"
                    )
                    pix.save(str(image_path))
                    try:
                        text = image_component.extract(image_path)
                    finally:
                        image_path.unlink(missing_ok=True)

                for chunk_index, chunk in enumerate(
                    _chunk_page(text)
                ):
                    chunks.append(
                        {
                            "text": chunk,
                            "page": page_number,
                            "page_chunk": chunk_index,
                        }
                    )
        finally:
            doc.close()

        return chunks

    if ext == ".docx":
        document = docx.Document(path)
        parts = [
            p.text.strip()
            for p in document.paragraphs
            if p.text.strip()
        ]

        # Tables often contain the exact approval limits, responsibilities,
        # thresholds and procedures users will ask the RAG system about.
        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip().replace("\n", " ")
                    for cell in row.cells
                ]
                row_text = " | ".join(
                    cell for cell in cells if cell
                )
                if row_text:
                    parts.append(row_text)

        text = "\n".join(parts)
        return [
            {
                "text": chunk,
                "section": "document",
            }
            for chunk in _chunk_page(text)
        ]

    if ext in {".txt", ".log", ".csv", ".json", ".py"}:
        text = path.read_text(
            encoding="utf-8",
            errors="ignore",
        )
        return [
            {"text": chunk}
            for chunk in _chunk_page(text)
        ]

    if ext in {".png", ".jpg", ".jpeg"}:
        text = image_component.extract(path)
        return [
            {"text": chunk}
            for chunk in _chunk_page(text)
        ]

    raise ValueError(
        f"Unsupported knowledge file type: {ext}"
    )


class KnowledgeIngestionService:
    """Explicit, persistent ingestion of trusted organization documents."""

    async def ingest_path(
        self,
        path: Path,
        *,
        source: str | None = None,
    ) -> dict[str, Any]:
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported knowledge file type: {path.suffix}"
            )

        chunks = await asyncio.to_thread(
            _extract_file_sync,
            path,
        )

        source_name = str(source or path.name).strip() or path.name

        count = await vector_store.add_chunks(
            chunks,
            source_name,
        )

        return {
            "status": "INGESTED",
            "source": source_name,
            "chunk_count": count,
            "chroma_collection": "company_knowledge",
        }

    async def ingest_directory(self) -> list[dict[str, Any]]:
        """Rebuild the knowledge index from the trusted directory.

        A directory rebuild is intentionally destructive to the existing
        Chroma collection. The directory is the authoritative source of
        organization knowledge, so stale chunks from older documents must be
        removed before the current files are indexed.
        """
        results: list[dict[str, Any]] = []
        directory = settings.KNOWLEDGE_DIR
        directory.mkdir(parents=True, exist_ok=True)

        paths = [
            path
            for path in sorted(directory.iterdir())
            if path.is_file()
            and path.suffix.lower() in SUPPORTED_EXTENSIONS
        ]

        await vector_store.reset_collection()

        for path in paths:
            results.append(
                await self.ingest_path(path)
            )

        return results

    async def search(
        self,
        query: str,
        top_k: int = 4,
    ) -> list[dict[str, Any]]:
        return await rag_engine.retrieve(
            query,
            top_k=top_k,
        )


knowledge_ingestion_service = KnowledgeIngestionService()
