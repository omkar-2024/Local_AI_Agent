import csv
import io
import re
from pathlib import Path

import docx
from openpyxl import Workbook
from pptx import Presentation

from app.core.config import settings


def _safe_output_path(filename: str) -> Path:
    """
    Return a file path strictly inside WORKSPACE_DIR.

    The LLM is allowed to choose a filename, but it must never be able to
    write outside the workbench workspace.
    """
    raw = Path(str(filename or "")).name
    raw = re.sub(r"[^A-Za-z0-9._-]+", "_", raw).strip("._")

    if not raw:
        raw = "generated_file"

    path = (settings.WORKSPACE_DIR / raw).resolve()
    workspace = settings.WORKSPACE_DIR.resolve()

    if path.parent != workspace:
        raise ValueError("Invalid output filename.")

    return path


def generate_word_document(filename: str, content: str) -> str:
    """Content uses '# ' for the title, '## ' for headings, plain lines for paragraphs."""
    path = _safe_output_path(filename)
    if path.suffix.lower() != ".docx":
        path = path.with_suffix(".docx")

    document = docx.Document()

    for line in str(content).splitlines():
        if line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif line.strip():
            document.add_paragraph(line.strip())

    document.save(path)
    return path.relative_to(settings.BASE_DIR).as_posix()


def generate_excel_document(filename: str, csv_data: str) -> str:
    """csv_data is plain CSV text; each row becomes a spreadsheet row."""
    path = _safe_output_path(filename)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")

    workbook = Workbook()
    sheet = workbook.active

    reader = csv.reader(io.StringIO(str(csv_data)))
    for row in reader:
        sheet.append(row)

    workbook.save(path)
    return path.relative_to(settings.BASE_DIR).as_posix()


def generate_pptx_document(filename: str, content: str) -> str:
    """Content uses '# ' to start a new slide title, '- ' lines beneath become bullet points."""
    path = _safe_output_path(filename)
    if path.suffix.lower() != ".pptx":
        path = path.with_suffix(".pptx")

    presentation = Presentation()
    bullet_layout = presentation.slide_layouts[1]
    slide = None
    body = None
    first_bullet = True

    for line in str(content).splitlines():
        if line.startswith("# "):
            slide = presentation.slides.add_slide(bullet_layout)
            slide.shapes.title.text = line[2:].strip()
            body = slide.placeholders[1].text_frame
            body.clear()
            first_bullet = True

        elif line.startswith("- ") and slide is not None:
            paragraph = (
                body.paragraphs[0]
                if first_bullet
                else body.add_paragraph()
            )
            paragraph.text = line[2:].strip()
            first_bullet = False

    presentation.save(path)
    return path.relative_to(settings.BASE_DIR).as_posix()
